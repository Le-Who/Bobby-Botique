"""Bound extracted Word text before storage and provider context assembly."""

import io
from types import SimpleNamespace

import pytest
from docx import Document

from app import document_processor as processor_module


@pytest.mark.parametrize("in_table", [False, True])
def test_word_extraction_bounds_paragraph_and_table_text(monkeypatch, in_table):
    document = Document()
    if in_table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Большой текст " * 100
    else:
        document.add_paragraph("Большой текст " * 100)
    data = io.BytesIO()
    document.save(data)
    data.seek(0)
    monkeypatch.setattr(processor_module, "MAX_DOCUMENT_TEXT_LENGTH", 100)

    result = processor_module.DocumentProcessor._process_word_sync(data)

    assert result["success"] is True
    assert len(result["content"]) <= 100
    assert result["text_length"] == len(result["content"])
    assert result["truncated"] is True


@pytest.mark.asyncio
async def test_pdf_extraction_reports_its_bounded_content(monkeypatch):
    monkeypatch.setattr(processor_module, "MAX_DOCUMENT_TEXT_LENGTH", 100)
    monkeypatch.setattr(
        processor_module.pypdf,
        "PdfReader",
        lambda _stream: SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "Текст " * 100)]),
    )
    processor = processor_module.DocumentProcessor()
    result = await processor._process_pdf_unified(b"%PDF synthetic bytes", "large.pdf")
    assert result["success"] is True
    assert len(result["content"]) <= 100
    assert result["text_length"] == len(result["content"])
    assert result["truncated"] is True


def test_small_word_document_preserves_complete_text():
    document = Document()
    document.add_paragraph("Первый абзац")
    document.add_paragraph("Второй абзац")
    data = io.BytesIO()
    document.save(data)
    data.seek(0)
    result = processor_module.DocumentProcessor._process_word_sync(data)
    assert result["content"] == "Первый абзац\n\nВторой абзац"
    assert result["truncated"] is False
